from __future__ import annotations

import csv
import gzip
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .util import now, read_json, write_json, adapter_sha
from .results import event_metrics
from .concurrency import read_concurrency

COLUMNS = ["task", "outcome", "score", "reward", "model", "framework_sha", "attempt", "sessionId",
           "elapsed_seconds", "agent_seconds", "model_rounds", "tool_calls", "input_tokens", "output_tokens",
           "skills_provided", "skills_discoverable", "skills_invoked", "skills_loaded", "failed_tests", "error", "evidence_complete",
           "analysis_status", "analysis_audit_status", "primary_attribution", "analysis_confidence", "detailed_analysis", "evidence",
           "execution_adapter_sha", "execution_concurrency"]


def usage(directory: Path, request_id: str | None) -> dict:
    totals, counts, seen = Counter(), Counter(), set()
    for file in sorted((directory / "nextagent-home" / "logs").glob("*")):
        if not file.name.endswith((".jsonl", ".jsonl.gz", ".ndjson", ".ndjson.gz")):
            continue
        opener = gzip.open if file.name.endswith(".gz") else open
        with opener(file, "rt", encoding="utf-8", errors="replace") as stream:
            for line in stream:
                try:
                    entry = json.loads(line)
                except ValueError:
                    continue
                if entry.get("event") != "model.payload.output_captured" or not request_id or entry.get("requestId") != request_id:
                    continue
                key = (entry.get("timestamp"), entry.get("requestId"), entry.get("stepId"))
                if key in seen:
                    continue
                seen.add(key)
                for name, value in (entry.get("modelOutput", {}).get("usage") or {}).items():
                    if name in {"inputTokens", "outputTokens"} and isinstance(value, (int, float)):
                        totals[name] += value
                        counts[name] += 1
    return {"input_tokens": totals["inputTokens"] if counts["inputTokens"] else None,
            "output_tokens": totals["outputTokens"] if counts["outputTokens"] else None,
            "usage_records": len(seen), "usage_field_counts": dict(counts)}


def analysis(directory: Path, result: dict) -> str:
    terminal = result.get("terminal", {})
    payload = terminal.get("payload", {})
    code = payload.get("code") or payload.get("guardReason")
    failed = result.get("verifier", {}).get("failed_tests", [])
    lines = [f"# {result['task']} / attempt {result['attempt']}", "", f"Outcome: {result.get('outcome')}",
             f"End-to-end score: {result.get('score')}; verifier reward: {result.get('reward')}",
             f"Model: {result.get('model')}; framework: {result.get('framework_sha')}", "", "## Evidence", "",
             "- result.json: phase, terminal, budgets, timings and collection status.",
             "- agent-facing/adaptation.json and adaptation.patch: exact path changes; verifier unchanged.",
             "- skill-catalog.json + skill-discovery.json: discoverability; trajectory.jsonl: actual Skill invocation and loading.",
             "- nextagent-home.tar: complete POSIX execution workspace, persistent state and framework logs; nextagent-home/logs/: readable log copy.",
             "- ctrf.json, reward.txt and verifier-stdout.txt: original scoring evidence.", "", "## Finding", ""]
    if code:
        lines.append(f"Observed terminal code: `{code}`. Terminal sequence: {terminal.get('sequence', 'unknown')}.")
    if result.get("error"):
        lines.append(result["error"])
    if failed:
        lines.append("Failed verifier cases: " + ", ".join(failed))
    if result.get("outcome") == "PASSED":
        lines.append("The agent completed and the original verifier awarded full reward.")
    elif code == "MODEL_DUPLICATION_DETECTED":
        lines.append("This code can originate from NextAgent's local duplication detector. It does not prove HTTP 429 or provider overload. Inspect model invocation and duplication details; no provider root cause is asserted here.")
    else:
        lines.append("Root cause is not established by this automatic summary. Distinguish task assertions, framework execution, model/provider evidence and harness/environment errors before assigning responsibility.")
    if result.get("artifact_errors"):
        lines.extend(["", "Evidence collection errors: " + "; ".join(result["artifact_errors"])])
    return "\n".join(lines) + "\n"


def row_for(directory: Path, result: dict) -> dict:
    result = dict(result)
    trajectory = directory / "trajectory.jsonl"
    if trajectory.exists():
        events = []
        for line in trajectory.read_text(encoding="utf-8").splitlines():
            try:
                events.append(json.loads(line))
            except ValueError:
                continue
        result.update(event_metrics(events))
    row = {name: result.get(name) for name in COLUMNS}
    row.update(usage(directory, result.get("requestId")))
    row["agent_seconds"] = result.get("phase_seconds", {}).get("agent")
    row["skills_provided"] = "; ".join(item["name"] for item in result.get("provided_skills", []))
    for key in ("skills_discoverable", "skills_invoked", "skills_loaded"):
        row[key] = "; ".join(result.get(key, []))
    row["failed_tests"] = "; ".join(result.get("verifier", {}).get("failed_tests", []))
    analysis = directory / "analysis" / "latest.json"
    if analysis.exists():
        details = read_json(analysis)
        row["analysis_status"] = details.get("status")
        row["analysis_audit_status"] = details.get("audit_status")
        row["primary_attribution"] = details.get("root_cause", {}).get("owner")
        row["analysis_confidence"] = details.get("root_cause", {}).get("confidence")
        row["detailed_analysis"] = details.get("report") or details.get("directory")
    row["evidence"] = str(directory.resolve())
    return row


def generate_report(directory: Path) -> dict:
    manifest = read_json(directory / "manifest.json")
    first, attempts = [], []
    for task in manifest["tasks"]:
        records = []
        for path in sorted((directory / "tasks" / task).glob("attempt-*/result.json")):
            result = read_json(path)
            row = row_for(path.parent, result)
            records.append(row)
            attempts.append(row)
            (path.parent / "failure-analysis.md").write_text(analysis(path.parent, result), encoding="utf-8")
        first.append(next((row for row in records if isinstance(row.get("score"), (int, float))), records[0] if records else
                          {"task": task, "outcome": "PENDING", "model": manifest["model"], "framework_sha": manifest["framework_sha"]}))
    counts = Counter(row.get("outcome", "RUNNING") for row in first)
    scored = [row["score"] for row in first if isinstance(row.get("score"), (int, float))]
    summary = {"generated_at": now(), "report_adapter_sha": adapter_sha(), "execution_adapter_sha": manifest.get("adapter_sha"),
               "runtime_adapter_sha": manifest.get("runtime_adapter_sha", manifest.get("adapter_sha")),
               "target_concurrency": read_concurrency(directory, manifest.get("jobs", 1)),
               "run_id": manifest["run_id"], "status": manifest["status"], "planned_tasks": len(first),
               "attempts": len(attempts), "first_attempt_outcomes": dict(counts), "scored_tasks": len(scored),
               "coverage": len(scored) / len(first) if first else 0,
               "full_suite_score": sum(scored) / len(first) if first and len(scored) == len(first) else None,
               "score_sum": sum(scored), "note": "First valid scored attempt is primary. A scored failure is never replaced by a later success. All invalid/interrupted attempts remain visible."}
    write_json(directory / "summary.json", summary)
    write_json(directory / "results.json", {"summary": summary, "first_attempts": first, "attempts": attempts})
    with (directory / "results.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(first)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "批次概览"
    for key, value in summary.items():
        overview.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value])
    for name, rows in (("首次有效结果", first), ("全部尝试", attempts)):
        sheet = workbook.create_sheet(name)
        sheet.append(COLUMNS)
        for row in rows:
            sheet.append([str(row.get(key))[:32000] if isinstance(row.get(key), (dict, list)) else row.get(key) for key in COLUMNS])
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244A68")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
        sheet.column_dimensions["A"].width = 36
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = cell.value[:32000]
                    # Treat task/model text as data, not Excel formulas.
                    if cell.value.startswith(("=", "+", "-", "@")):
                        cell.value = "'" + cell.value
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    overview.column_dimensions["A"].width = 32
    overview.column_dimensions["B"].width = 100
    workbook.save(directory / "results.xlsx")
    lines = ["# NextAgent SkillsBench results", "", f"Run: {manifest['run_id']}", f"Model: {manifest['model']}",
             f"Framework: {manifest['framework_sha']}", f"Status: {manifest['status']}",
             f"Scored coverage: {len(scored)}/{len(first)}", f"Full-suite score: {summary['full_suite_score']}", "",
             "| Task | First outcome | Score | Verifier reward | Evidence |", "|---|---|---:|---:|---|"]
    for row in first:
        lines.append(f"| {row['task']} | {row.get('outcome')} | {row.get('score')} | {row.get('reward')} | {row.get('evidence', '')} |")
    (directory / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return summary
