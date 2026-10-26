from __future__ import annotations

import json
import tempfile
import tarfile
import unittest
from pathlib import Path

from nextagent_bench.artifacts import import_framework, inspect_archive
from nextagent_bench.paths import predict_workdir
from nextagent_bench.report import generate_report
from nextagent_bench.runner import check_resume, create_run, poll_run
from nextagent_bench.settings import load_settings
from nextagent_bench.util import read_json, write_json


class EventAPI:
    def __init__(self):
        self.paths = []

    def rpc(self, method, path, body=None):
        self.paths.append(path)
        if "afterSequence=0&" in path:
            return {"events": [{"eventId": str(i), "sequence": i, "eventType": "LLM_CONTENT_DELTA", "payload": {}} for i in range(1, 1001)]}
        return {"events": [{"eventId": "1000", "sequence": 1000, "eventType": "LLM_CONTENT_DELTA", "payload": {}},
                           {"eventId": "1001", "sequence": 1001, "eventType": "GUARD_BLOCKED", "payload": {"guardReason": "DENIED"}}]}


class PublicEventContractTests(unittest.TestCase):
    def test_terminal_after_first_page_and_duplicate_events_are_handled(self):
        api = EventAPI()
        with tempfile.TemporaryDirectory() as folder:
            trajectory = Path(folder) / "events.jsonl"
            result = poll_run(api, "session", "run", "request", 10, trajectory)
            events = [json.loads(line) for line in trajectory.read_text().splitlines()]
        self.assertEqual(result["status"], "GUARD_BLOCKED")
        self.assertEqual(len(events), 1001)
        self.assertIn("afterSequence=1000&", api.paths[-1])

    def test_stop_is_observed_without_waiting_for_a_new_event(self):
        with tempfile.TemporaryDirectory() as folder:
            result = poll_run(EventAPI(), "s", "r", "q", 10, Path(folder) / "events.jsonl", lambda: True)
        self.assertEqual(result["status"], "INTERRUPTED")


class SnapshotTests(unittest.TestCase):
    def test_release_archive_cannot_escape_extraction_root(self):
        import io
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "malicious.tar.gz"
            with tarfile.open(path, "w:gz") as archive:
                item = tarfile.TarInfo("../../outside")
                item.size = 1
                archive.addfile(item, io.BytesIO(b"x"))
            with self.assertRaisesRegex(ValueError, "Unsafe archive"):
                inspect_archive(path)

    def test_framework_import_drops_root_runtime_state_but_preserves_nested_data(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source"
            (source / "node_modules/pkg/data").mkdir(parents=True)
            (source / "node_modules/pkg/data/required.json").write_text("{}")
            (source / "logs").mkdir()
            (source / "logs/private.txt").write_text("runtime state")
            write_json(source / "package.json", {"dependencies": {"@nextagent/agent-app": "1"}})
            write_json(source / "candidate-manifest.json", {"platform": "linux", "arch": "x64"})
            artifact = import_framework(root / "project", source)
            bundle = root / "project/artifacts" / artifact["sha256"] / "bundle.tar.gz"
            with tarfile.open(bundle) as archive:
                names = {name.removeprefix("./") for name in archive.getnames()}
                self.assertIn("node_modules/pkg/data/required.json", names)
                self.assertNotIn("logs/private.txt", names)

    def test_final_stage_workdir_and_environment_are_resolved(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "Dockerfile"
            path.write_text("FROM x AS builder\nWORKDIR /wrong\nFROM y\nENV WORKSPACE=/app/workspace \\\n OTHER=value\nWORKDIR ${WORKSPACE}\nWORKDIR subdir\n")
            self.assertEqual(predict_workdir(path), "/app/workspace/subdir")

    def test_resume_detects_changed_config_contents(self):
        actual_project = Path(__file__).resolve().parents[1]
        config = load_settings(actual_project)
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            run = create_run(root, config, {"sha256": "a"}, {"sha256": "b"}, ["task"], 1)
            self.assertEqual(check_resume(run)["model"], config["model"])
            agent = read_json(run / "config/agent.yaml")
            agent["runtimeSettings"]["guardrailEnabled"] = True
            write_json(run / "config/agent.yaml", agent)
            with self.assertRaisesRegex(ValueError, "configuration"):
                check_resume(run)


class ReportingTests(unittest.TestCase):
    def test_retry_success_does_not_replace_scored_failure(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            write_json(root / "manifest.json", {"run_id": "r", "status": "FINISHED", "model": "qwen3.8-27b", "framework_sha": "a", "tasks": ["task", "pending"]})
            for number, score in [(1, 0), (2, 1)]:
                write_json(root / f"tasks/task/attempt-{number:03}/result.json", {"task": "task", "attempt": number,
                           "score": score, "reward": score, "outcome": "TASK_FAILED" if score == 0 else "PASSED", "model": "qwen3.8-27b"})
            summary = generate_report(root)
            self.assertEqual(summary["score_sum"], 0)
            self.assertIsNone(summary["full_suite_score"])
            book = load_workbook(root / "results.xlsx", read_only=True)
            sheet = book["首次有效结果"]
            self.assertEqual(sheet.cell(2, 3).value, 0)
            self.assertEqual(sheet.cell(2, 5).value, "qwen3.8-27b")
            book.close()


if __name__ == "__main__":
    unittest.main()
